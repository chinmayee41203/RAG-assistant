import io
from pathlib import Path
from typing import List

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from config import CHUNK_SIZE, CHUNK_OVERLAP


def load_txt(file_bytes: bytes, filename: str) -> List[Document]:
    text = file_bytes.decode("utf-8", errors="replace")
    return [Document(page_content=text, metadata={"source": filename, "page": 1})]


def load_pdf(file_bytes: bytes, filename: str) -> List[Document]:
    import pypdf

    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    docs = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(Document(page_content=text, metadata={"source": filename, "page": i}))
    return docs


def load_docx(file_bytes: bytes, filename: str) -> List[Document]:
    import docx

    doc = docx.Document(io.BytesIO(file_bytes))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [Document(page_content=text, metadata={"source": filename, "page": 1})]


def load_document(file_bytes: bytes, filename: str) -> List[Document]:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return load_pdf(file_bytes, filename)
    elif ext == ".docx":
        return load_docx(file_bytes, filename)
    elif ext == ".txt":
        return load_txt(file_bytes, filename)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def split_documents(docs: List[Document]) -> List[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(docs)
    for i, chunk in enumerate(chunks):
        chunk.metadata["chunk_index"] = i
    return chunks


def process_uploaded_file(file_bytes: bytes, filename: str) -> List[Document]:
    raw_docs = load_document(file_bytes, filename)
    return split_documents(raw_docs)
