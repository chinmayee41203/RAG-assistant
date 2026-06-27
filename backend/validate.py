"""
Run this script to verify the full RAG pipeline works end-to-end before launching the app.
Usage: python validate.py
"""

import sys
import textwrap

PASS = "\033[92m[PASS]\033[0m"
FAIL = "\033[91m[FAIL]\033[0m"
INFO = "\033[94m[INFO]\033[0m"


def check(label: str, fn):
    try:
        result = fn()
        print(f"{PASS} {label}" + (f" — {result}" if result else ""))
        return True
    except Exception as e:
        print(f"{FAIL} {label} — {e}")
        return False


def test_imports():
    import langchain          # noqa: F401
    import chromadb           # noqa: F401
    import sentence_transformers  # noqa: F401
    import pypdf              # noqa: F401
    import docx               # noqa: F401
    import streamlit          # noqa: F401
    return "all packages importable"


def test_config():
    import config
    assert config.CHUNK_SIZE > 0
    assert config.CHUNK_OVERLAP >= 0
    assert config.LLM_PROVIDER in ("ollama", "openai")
    return f"provider={config.LLM_PROVIDER}, model={config.OLLAMA_MODEL if config.LLM_PROVIDER == 'ollama' else config.OPENAI_MODEL}"


def test_txt_processing():
    from document_processor import process_uploaded_file
    sample = b"This is a test document.\nIt has multiple sentences for chunking validation."
    chunks = process_uploaded_file(sample, "test.txt")
    assert len(chunks) >= 1
    assert chunks[0].metadata["source"] == "test.txt"
    return f"{len(chunks)} chunk(s)"


def test_pdf_processing():
    import io
    import pypdf
    from pypdf import PdfWriter
    from document_processor import process_uploaded_file

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    pdf_bytes = buf.getvalue()
    chunks = process_uploaded_file(pdf_bytes, "test.pdf")
    return f"{len(chunks)} chunk(s) (blank page expected 0 text chunks)"


def test_docx_processing():
    import io
    import docx
    from document_processor import process_uploaded_file

    doc = docx.Document()
    doc.add_paragraph("Hello from a DOCX file. This is a validation paragraph.")
    buf = io.BytesIO()
    doc.save(buf)
    chunks = process_uploaded_file(buf.getvalue(), "test.docx")
    assert len(chunks) >= 1
    return f"{len(chunks)} chunk(s)"


def test_embeddings():
    from rag_pipeline import get_embeddings
    emb = get_embeddings()
    result = emb.embed_query("hello world")
    assert len(result) > 0
    return f"vector dim={len(result)}"


def test_chroma_ingest_and_retrieve():
    from langchain_core.documents import Document
    from rag_pipeline import get_vector_store, clear_vector_store

    clear_vector_store()
    store = get_vector_store()

    docs = [
        Document(
            page_content="The Eiffel Tower is located in Paris, France.",
            metadata={"source": "_validate_test.txt", "page": 1, "chunk_index": 0},
        ),
        Document(
            page_content="Python is a high-level programming language.",
            metadata={"source": "_validate_test.txt", "page": 1, "chunk_index": 1},
        ),
    ]
    store.add_documents(docs)

    results = store.similarity_search("Where is the Eiffel Tower?", k=1)
    assert len(results) >= 1
    assert "Paris" in results[0].page_content

    clear_vector_store()
    return "ingest + retrieve verified"


def test_ollama_connection():
    import config
    if config.LLM_PROVIDER != "ollama":
        return "skipped (provider is not ollama)"
    import requests
    resp = requests.get(f"{config.OLLAMA_BASE_URL}/api/tags", timeout=5)
    resp.raise_for_status()
    models = [m["name"] for m in resp.json().get("models", [])]
    if config.OLLAMA_MODEL not in models and not any(config.OLLAMA_MODEL in m for m in models):
        raise RuntimeError(
            f"Model '{config.OLLAMA_MODEL}' not found in Ollama. "
            f"Available: {models}. Run: ollama pull {config.OLLAMA_MODEL}"
        )
    return f"Ollama reachable, model '{config.OLLAMA_MODEL}' available"


def test_full_rag_query():
    from langchain_core.documents import Document
    from rag_pipeline import add_documents, query_rag, clear_vector_store

    clear_vector_store()
    docs = [
        Document(
            page_content="The RAG validation document states that the secret number is 4291.",
            metadata={"source": "_validate_rag.txt", "page": 1, "chunk_index": 0},
        )
    ]
    add_documents(docs)
    result = query_rag("What is the secret number mentioned in the document?")
    clear_vector_store()

    assert "answer" in result
    assert "sources" in result
    assert len(result["answer"]) > 0
    snippet = result["sources"][0]["snippet"] if result["sources"] else ""
    return f"answer received ({len(result['answer'])} chars), source: {result['sources'][0]['source'] if result['sources'] else 'none'}"


if __name__ == "__main__":
    print("\n" + "=" * 55)
    print("  RAG 2.0 — Validation Suite")
    print("=" * 55 + "\n")

    tests = [
        ("Package imports",            test_imports),
        ("Config loading",             test_config),
        ("TXT document processing",    test_txt_processing),
        ("PDF document processing",    test_pdf_processing),
        ("DOCX document processing",   test_docx_processing),
        ("Embedding model",            test_embeddings),
        ("ChromaDB ingest & retrieve", test_chroma_ingest_and_retrieve),
        ("Ollama connectivity",        test_ollama_connection),
        ("Full RAG query (LLM)",       test_full_rag_query),
    ]

    results = [check(label, fn) for label, fn in tests]

    passed = sum(results)
    total = len(results)
    print(f"\n{'=' * 55}")
    print(f"  Results: {passed}/{total} passed")
    print("=" * 55 + "\n")

    if passed < total:
        sys.exit(1)
